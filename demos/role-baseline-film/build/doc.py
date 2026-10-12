"""Compile out/补充说明.txt into a plain Word file and a plain PDF: headings, paragraphs and one table, nothing else.

  python3 build/doc.py        → out/补充说明.docx, out/补充说明.pdf

The txt stays the source. Structure is read from it: the first line is the title, lines that open with
"一、" … are section headings, the aligned block under "单位：" is the table, and a paragraph that opens
with a short label and "：" in the last section keeps the label in bold.
Word comes from python-docx, the PDF from headless Chromium printing the same structure as HTML.
"""
import html, re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'out' / '补充说明.txt'
HEAD = re.compile(r'^[一二三四五六七八九十]+、')
LABEL = re.compile(r'^([^：，。；]{2,5})：(.+)$')


def parse(text):
    """→ [('title', s) | ('h', s) | ('p', s) | ('cap', s) | ('item', label, s) | ('table', header, rows)]
    A 'cap' is the paragraph right above the table; it stays on the table's page."""
    lines = text.splitlines()
    out, i, last = [('title', lines[0].strip())], 1, False
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if HEAD.match(ln):
            out.append(('h', ln.strip())); last = ln.startswith('五、'); i += 1; continue
        if ln.startswith(' ') and len(ln.split()) == 2:          # table header: two column names, indented
            header, rows = ln.split(), []
            i += 1
            while i < len(lines) and len(lines[i].split()) == 3:
                rows.append(lines[i].split()); i += 1
            if out[-1][0] == 'p': out[-1] = ('cap', out[-1][1])
            out.append(('table', header, rows)); continue
        m = LABEL.match(ln) if last else None
        out.append(('item', m.group(1), m.group(2)) if m else ('p', ln.strip()))
        i += 1
    return out


def east_asia(style, latin, cjk):
    style.font.name = latin
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), cjk)


def word(blocks, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.8); sec.top_margin = sec.bottom_margin = Cm(2.54)
    normal = doc.styles['Normal']
    east_asia(normal, 'Times New Roman', '宋体')
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.5
    for name, size, before in (('Heading 1', 16, 0), ('Heading 2', 13, 14)):
        st = doc.styles[name]
        east_asia(st, 'Arial', '黑体')
        st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(8)
    for b in blocks:
        if b[0] == 'title':
            doc.add_heading(b[1], level=1)
        elif b[0] == 'h':
            doc.add_heading(b[1], level=2)
        elif b[0] == 'p':
            doc.add_paragraph(b[1])
        elif b[0] == 'cap':
            doc.add_paragraph(b[1]).paragraph_format.keep_with_next = True
        elif b[0] == 'item':
            p = doc.add_paragraph()
            p.add_run(b[1] + '：').bold = True
            p.add_run(b[2])
        else:
            _, header, rows = b
            t = doc.add_table(rows=1 + len(rows), cols=3)
            t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.LEFT
            borders = OxmlElement('w:tblBorders')       # written out, so viewers that ignore the table style still draw the grid
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                el = OxmlElement(f'w:{edge}')
                for k, v in (('val', 'single'), ('sz', '4'), ('space', '0'), ('color', '000000')): el.set(qn(f'w:{k}'), v)
                borders.append(el)
            t._tbl.tblPr.append(borders)
            for r, cells in enumerate([['项目'] + header] + rows):
                for c, s in enumerate(cells):
                    cell = t.cell(r, c)
                    cell.width = Cm((4.2, 3.4, 3.4)[c])
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(2); par.paragraph_format.line_spacing = 1.2
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT if c else WD_ALIGN_PARAGRAPH.LEFT
                    par.paragraph_format.keep_with_next = r < len(rows)
                    par.add_run(s).bold = r == 0
            doc.add_paragraph()
    doc.save(path)


CSS = """
@page { size: A4; margin: 25.4mm 28mm; }
body { font: 11pt/1.75 "Songti SC", "SimSun", "Times New Roman", serif; color: #000; margin: 0; }
h1, h2 { font-family: "PingFang SC", "Heiti SC", "SimHei", Arial, sans-serif; font-weight: 600; break-after: avoid; }
h1 { font-size: 16pt; margin: 0 0 10pt; }
h2 { font-size: 13pt; margin: 18pt 0 6pt; }
p { margin: 0 0 8pt; text-align: justify; }
table { border-collapse: collapse; margin: 4pt 0 12pt; }
.keep { break-inside: avoid; }
th, td { border: 0.5pt solid #000; padding: 3pt 12pt; font-weight: 400; text-align: right; }
th { font-weight: 600; }
th:first-child, td:first-child { text-align: left; }
"""


def pdf(blocks, path):
    e = html.escape
    body = []
    for b in blocks:
        if b[0] == 'title': body.append(f'<h1>{e(b[1])}</h1>')
        elif b[0] == 'h': body.append(f'<h2>{e(b[1])}</h2>')
        elif b[0] == 'p': body.append(f'<p>{e(b[1])}</p>')
        elif b[0] == 'cap': body.append(f'<div class="keep"><p>{e(b[1])}</p>')
        elif b[0] == 'item': body.append(f'<p><b>{e(b[1])}：</b>{e(b[2])}</p>')
        else:
            _, header, rows = b
            tr = lambda cells, tag: '<tr>' + ''.join(f'<{tag}>{e(c)}</{tag}>' for c in cells) + '</tr>'
            body.append('<table>' + tr(['项目'] + header, 'th') + ''.join(tr(r, 'td') for r in rows) + '</table></div>')
    page_html = f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>{e(blocks[0][1])}</title><style>{CSS}</style>{"".join(body)}</html>'
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        page = browser.new_page()
        page.set_content(page_html)
        page.pdf(path=str(path), format='A4', prefer_css_page_size=True)
        browser.close()


def main():
    blocks = parse(SRC.read_text())
    kinds = [b[0] for b in blocks]
    assert kinds.count('table') == 1 and kinds.count('h') == 5, kinds
    word(blocks, SRC.with_suffix('.docx'))
    pdf(blocks, SRC.with_suffix('.pdf'))
    assert kinds[kinds.index('table') - 1] == 'cap', kinds
    print({k: kinds.count(k) for k in ('h', 'p', 'cap', 'item', 'table')})


if __name__ == '__main__':
    main()
